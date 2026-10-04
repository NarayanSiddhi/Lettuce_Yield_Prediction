from pathlib import Path
from datetime import datetime

from docx import Document


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def main() -> None:
    project_root = Path(__file__).resolve().parent
    out_path = project_root / "Hydroponic_Lettuce_Project_Progress_Report.docx"

    doc = Document()
    doc.add_heading("Hydroponic Lettuce Project Progress Report", 0)
    doc.add_paragraph(f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    doc.add_heading("1. Project Objective", level=1)
    doc.add_paragraph(
        "Build a research-focused machine learning pipeline that links hydroponic environmental "
        "and nutrient time-series data with lettuce growth outcomes, and includes image-derived "
        "signals from daily timelapse captures."
    )

    doc.add_heading("2. Source Data and Files", level=1)
    add_bullets(
        doc,
        [
            "Ambient_nutrient_data.xlsx: Environmental and nutrient sensor streams plus actuator logs.",
            "Lettuce_FW_EC_Tracker_v3_2_.xlsx: Fresh-weight measurements and checkpoint growth labels.",
            "Timelapse_growtent/Trial-1-Camera Capture...: Daily image folders (16 images/day).",
            "Project_requirements.docx and Methodology_Draft.docx: Project scope and experiment context.",
        ],
    )

    doc.add_heading("3. Workflow Completed (Step-by-Step)", level=1)

    doc.add_heading("3.1 Dataset Discovery and Requirement Alignment", level=2)
    add_bullets(
        doc,
        [
            "Verified project structure and validated relevant files/folders.",
            "Extracted objective and deliverables from Project_requirements.docx.",
            "Confirmed this is an insight + prediction + decision-support workflow (not deployment automation).",
        ],
    )

    doc.add_heading("3.2 Image Pipeline Design and Cropping Decisions", level=2)
    add_bullets(
        doc,
        [
            "Initial center/auto-green cropping was tested but did not consistently isolate a single plant.",
            "Switched to a fixed single-plant ROI approach across all days for consistency.",
            "Final fixed coordinates selected by user and applied globally: x1=1721, y1=673, x2=3418, y2=2002.",
            "Best image per day selected from 16 captures with visibility and quality preference.",
        ],
    )

    doc.add_heading("3.3 Scripts Created/Updated for Image Processing", level=2)
    add_bullets(
        doc,
        [
            "make_daily_crops_trial1.py (initial exploratory crop generation).",
            "make_single_plant_crops_trial1.py (fixed single-plant ROI + best-of-16 frame selection).",
            "get_black_box_coords.py (interactive coordinate capture helper).",
        ],
    )

    doc.add_heading("3.4 Image Outputs Produced", level=2)
    add_bullets(
        doc,
        [
            "crops/trial1_singleplant/: 28 final cropped daily images (one/day).",
            "crops/trial1_singleplant_preview/: boxed source-frame previews for QA.",
            "image_daily_index_trial1_singleplant.csv: date-wise mapping of selected source image and ROI metadata.",
        ],
    )

    doc.add_heading("3.5 Image Feature Engineering (Step 2)", level=2)
    add_bullets(
        doc,
        [
            "Created script: extract_image_features_trial1_singleplant.py.",
            "Generated output: image_features_trial1_singleplant.csv (28 rows, one/day).",
            "Features extracted: green_ratio, vegetation_ratio, canopy_area_px, mean_h, mean_s, mean_v, color_std, blur_score, texture_edge_density, texture_laplacian_mean_abs.",
            "Improvement applied: added vegetation_ratio to capture red/purple lettuce canopy better than green_ratio alone.",
        ],
    )

    doc.add_heading("3.6 Sensor and Label Integration (Step 3)", level=2)
    add_bullets(
        doc,
        [
            "Installed dependencies: openpyxl (Excel parsing), python-docx (report generation).",
            "Created script: build_master_dataset_trial1.py.",
            "Parsed and aggregated Environmental Data and Nutrient Data to daily features (means/std/min/max + actuator sums/event counts).",
            "Merged daily sensor features with daily image features on date.",
            "Built 4-day checkpoint modeling table aligned to median fresh-weight labels.",
        ],
    )

    doc.add_heading("4. Final CSV Outputs Created", level=1)
    add_bullets(
        doc,
        [
            "sensor_features_daily_trial1.csv (29 rows): daily sensor/actuator aggregates.",
            "merged_daily_features_trial1.csv (28 rows): image + sensor daily merged table.",
            "master_checkpoint_dataset_trial1.csv (8 rows): checkpoint-window modeling dataset with target_median_fw_g.",
            "image_features_trial1_singleplant.csv (28 rows): final daily image feature table.",
        ],
    )

    doc.add_heading("5. Important Notes and Assumptions", level=1)
    add_bullets(
        doc,
        [
            "Day-28 checkpoint window includes only 3 days of available image timeline data; this is recorded as window_days_count=3.",
            "Current model design uses image-derived numeric features in tabular ML, not raw-image deep learning.",
            "This pipeline is designed for robust baseline research analysis and can be extended later (multi-image/day summaries, CNNs, etc.).",
        ],
    )

    doc.add_heading("6. Recommended Next Step", level=1)
    doc.add_paragraph(
        "Train baseline predictive models (e.g., Linear Regression, Random Forest, XGBoost) using "
        "master_checkpoint_dataset_trial1.csv with time-aware validation, then produce feature-importance "
        "and actionable control recommendations."
    )

    doc.save(out_path)
    print(f"Report created: {out_path}")


if __name__ == "__main__":
    main()

